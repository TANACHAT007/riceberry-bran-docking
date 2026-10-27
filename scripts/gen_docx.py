"""Write abstract.json into the official INMU50 templates (keeps page setup + styles, replaces body)."""
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
A = json.loads((ROOT / "abstract.json").read_text())
OUT = ROOT / "submission"; OUT.mkdir(exist_ok=True)


def blank(path):
    d = Document(path)
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    return d


def run(p, text, font, size, bold=False, italic=False, underline=False, sup=False):
    r = p.add_run(text)
    r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.underline = underline
    r.font.superscript = sup
    rpr = r._element.get_or_add_rPr(); fonts = rpr.find(qn("w:rFonts"))
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), font)
    if size:
        szcs = rpr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(size * 2))}); rpr.append(szcs)
    return r


def para(d, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
    p = d.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    return p


def block(d, L, font, title_size, body_size, small_size, upper):
    p = para(d, WD_ALIGN_PARAGRAPH.CENTER, 8)
    run(p, L["title"].upper().replace("Α", "α") if upper else L["title"], font, title_size, bold=True)
    p = para(d, WD_ALIGN_PARAGRAPH.CENTER, 0)
    for i, au in enumerate(L["authors"]):
        if i: run(p, ", ", font, body_size)
        run(p, au["name"], font, body_size, underline=au.get("presenter", False))
        run(p, str(au["aff"]) + ("*" if au.get("corresponding") else ""), font, body_size, sup=True)
    for i, af in enumerate(L["affiliations"], 1):
        p = para(d, WD_ALIGN_PARAGRAPH.CENTER, 0)
        run(p, str(i), font, body_size, italic=True, sup=True); run(p, " " + af, font, body_size, italic=True)
    para(d, after=0)
    p = para(d)
    for i, (head, text) in enumerate(L["sections"]):
        run(p, ("" if i == 0 else " ") + head + ": ", font, body_size, bold=True); run(p, text, font, body_size)
    p = para(d, WD_ALIGN_PARAGRAPH.LEFT, 2)
    run(p, L["kw_label"] + ": ", font, small_size, bold=True); run(p, ", ".join(L["keywords"]), font, small_size)
    p = para(d, WD_ALIGN_PARAGRAPH.LEFT, 2)
    run(p, "*Corresponding author: ", font, small_size, bold=True); run(p, A["email"], font, small_size)


def words(L):
    return sum(len(t.split()) for _, t in L["sections"])


if __name__ == "__main__":
    en, th = A["en"], A["th"]
    d = blank(ROOT / "templates" / "en.docx")
    block(d, en, "Arial", 12, 11, 11, True)
    d.save(OUT / f"Abstract_{A['file_name']}_EN.docx")
    d = blank(ROOT / "templates" / "th.docx")
    block(d, th, "TH Sarabun New", 18, 16, 14, False)
    d.add_page_break()
    block(d, en, "Arial", 12, 11, 11, True)
    d.save(OUT / f"Abstract_{A['file_name']}.docx")
    print("EN body words:", words(en))
