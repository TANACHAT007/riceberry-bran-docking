"""Fill the official INMU50 poster template (80 x 142 cm) with results from summary.json."""
import json, sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
S = json.loads((ROOT / "summary.json").read_text())
A = json.loads((ROOT / "abstract.json").read_text())
FIG = ROOT / "figures"
NAVY, BERRY, GOLD, INK, SOFT, PALE = (RGBColor(0x1D, 0x3A, 0x8A), RGBColor(0x7A, 0x3F, 0xD1), RGBColor(0xC2, 0x8A, 0x00),
                                      RGBColor(0x22, 0x26, 0x33), RGBColor(0xEE, 0xF2, 0xFB), RGBColor(0xF4, 0xEE, 0xFF))
LEAF = RGBColor(0x16, 0x8A, 0x55)
FONT = "Sarabun"
X0, Y0, W = 7.0, 22.0, 66.4          # inside the template's content frame (6.1..74.3 cm, 21..123 cm)
COL = (W - 2.0) / 2


def text(slide, x, y, w, h, paras, size=26, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0.45):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Cm(0.2))
    for i, p in enumerate(paras):
        p = p if isinstance(p, list) else [(p, {})]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.space_after = Pt(size * space); para.line_spacing = 1.07
        for t, st in p:
            r = para.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(st.get("size", size)); f.bold = st.get("bold", False)
            f.italic = st.get("italic", False); f.color.rgb = st.get("color", color)
    return tb


def heading(slide, x, y, w, label, color=NAVY):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(2.6))
    s.adjustments[0] = 0.5; s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background(); s.shadow.inherit = False
    tf = s.text_frame; tf.margin_left = Cm(1.0); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; r = p.add_run(); r.text = label
    r.font.name = FONT; r.font.size = Pt(46); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    return y + 3.2


def panel(slide, x, y, w, h, fill=SOFT, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    s.adjustments[0] = 0.06; s.fill.solid(); s.fill.fore_color.rgb = fill; s.shadow.inherit = False
    if line:
        s.line.color.rgb = line; s.line.width = Pt(2)
    else:
        s.line.fill.background()
    return s


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w) if w else None, Cm(h) if h else None)


def B(t, **k): return (t, {"bold": True, **k})
def N(t, **k): return (t, k)
def res(r): return r[0] + r[1:3].lower() + r[3:]


def main(url=""):
    rows = {r["ligand"]: r for r in S["rows"]}
    v = S["validation"]
    acb, c3g, p3g = rows["acarbose"], rows["cyanidin-3-O-glucoside"], rows["peonidin-3-O-glucoside"]
    cf, q = rows["cycloartenyl ferulate"], rows["quercetin"]
    xt = min(r["affinity"] for r in v["MGAM"]["redock"])      # crystal-conformer acarbose on MGAM
    d_mgam = c3g["MGAM"]["mean"] - xt                          # ΔΔG vs acarbose, MGAM
    d_hpa = c3g["HPA"]["mean"] - acb["HPA"]["mean"]            # ΔΔG vs acarbose, HPA
    prs = Presentation(ROOT / "templates" / "poster.pptx"); sl = prs.slides[0]

    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("TITLE"):
            tf = sh.text_frame; tf.paragraphs[0].runs[0].text = A["poster_title"].upper().replace("Α", "α")
            for r in tf.paragraphs[0].runs[1:]: r.text = ""
            for p in tf.paragraphs[1:]:
                for r in p.runs: r.text = ""
        elif sh.has_text_frame and sh.text_frame.text.startswith("Firstname"):
            en = A["en"]; ps = sh.text_frame.paragraphs
            one = len(en["affiliations"]) == 1  # no superscript needed for a single affiliation
            ps[0].runs[0].text = ", ".join(a["name"] + ("" if one else str(a["aff"])) + ("*" if a.get("corresponding") else "") for a in en["authors"])
            for r in ps[0].runs[1:]: r.text = ""
            for i, p in enumerate(ps[1:]):
                if p.runs:
                    p.runs[0].text = ((en["affiliations"][i] if one else f"{i + 1} {en['affiliations'][i]}") if i < len(en["affiliations"]) else "")
                    for r in p.runs[1:]: r.text = ""
        elif sh.has_text_frame and sh.text_frame.text.startswith("*Corresponding"):
            sh.text_frame.paragraphs[0].runs[0].text = f"*Corresponding author: {A['corresponding_name']} ({A['email']})"
            for r in sh.text_frame.paragraphs[0].runs[1:]: r.text = ""

    # ============ Row 1: the problem | the question ============
    y = 22.0
    xr = X0 + COL + 2.0
    yb = heading(sl, X0, y, COL, "A BY-PRODUCT WORTH MINING")
    text(sl, X0, yb, COL, 13, [
        [N("Starch becomes glucose via pancreatic "), B("α-amylase"), N(" and intestinal "), B("α-glucosidase"),
         N(". Acarbose slows both, but is poorly tolerated (flatulence 74% vs 29% on placebo, US label).")],
        [B("Riceberry bran", color=BERRY), N(" is a low-value by-product of Thai rice milling. Pigmented-rice anthocyanins are "
         "already known "), B("from in vitro work"), N(" to inhibit both enzymes (Akkarachiyasit et al., 2010).")],
        [N("Bran carries many candidates at once, and a food ingredient must also "), B("dissolve where the enzymes are"),
         N(". These have never been compared on the human enzymes.")],
    ], size=27)
    yb2 = heading(sl, xr, y, COL, "THE QUESTION", BERRY)
    text(sl, xr, yb2, COL, 8.0, [
        [N("Riceberry bran, a "), B("low-value by-product", color=BERRY), N(" of Thai rice mills, carries anthocyanins, γ-oryzanol, vitamin E and phenolic acids.")],
        [B("Which fraction should a food project extract and assay first?"), N(" That is the question this screen answers.")]], size=27)
    panel(sl, xr, yb2 + 8.4, COL, 4.8, PALE)
    text(sl, xr + 0.8, yb2 + 8.8, COL - 1.6, 4.0, [
        [B("11 bran constituents ranked on both human enzymes, scoring predicted binding ", color=BERRY), B("and solubility together", color=BERRY)]], size=29)

    # ============ Row 2: methods ============
    y = 37.4
    yb = heading(sl, X0, y, W, "MATERIALS AND METHODS")
    steps = [("1  Targets", "Human MGAM (PDB 2QMJ, 1.9 Å) and pancreatic α-amylase (PDB 3BAJ, 2.1 Å), both crystallised with acarbose"),
             ("2  Protein prep", "PDBFixer: waters and ions removed, missing atoms and hydrogens added at pH 7.4; Meeko PDBQT"),
             ("3  Ligands", "12 structures from PubChem; RDKit ETKDG + MMFF94 optimisation; Meeko"),
             ("4  Docking", "AutoDock Vina 1.2.7 in the acarbose site, exhaustiveness 16, 3 independent seeds (78 runs)"),
             ("5  Analysis", "ΔΔG vs acarbose on each enzyme; contact residues ≤ 4 Å, H-bonds ≤ 3.5 Å, logP, ligand efficiency")]
    bw = (W - 4 * 0.7) / 5
    for i, (h, b) in enumerate(steps):
        bx = X0 + i * (bw + 0.7)
        panel(sl, bx, yb, bw, 9.6)
        text(sl, bx + 0.5, yb + 0.5, bw - 1.0, 8.8, [[B(h, color=NAVY, size=29)], [N(b, size=24)]])
    text(sl, X0, yb + 9.9, W, 2.4, [[B("Protocol validation: "), N(
        f"acarbose re-docked from a randomised start returned to its MGAM crystal pose in every run "
        f"(RMSD {v['MGAM']['rmsd']}–{max(r['rmsd'] for r in v['MGAM']['redock'])} Å, n = 3).")]], size=25)

    # ============ Row 3: the selectivity map (hero) ============
    y = 52.6
    yb = heading(sl, X0, y, W, "RESULTS  ·  RANKING THE BRAN")
    wl = 30.0; xr3 = X0 + wl + 1.4; wr = W - wl - 1.4
    picture(sl, FIG / "selectivity.png", X0, yb, w=wl)
    text(sl, X0, yb + 25.0, wl, 2.8, [[B("Fig. 1 "), N("Each constituent placed against acarbose on both enzymes (ΔΔG, mean of 3 seeds). "
         "Circles dissolve in water, squares do not. The shaded area marks where a usable food candidate would sit.")]], size=24)
    picture(sl, FIG / "pose_MGAM_cyanidin_3_O_glucoside.png", xr3, yb, w=wr / 2 - 0.3)
    picture(sl, FIG / "pose_HPA_cyanidin_3_O_glucoside.png", xr3 + wr / 2 + 0.3, yb, w=wr / 2 - 0.3)
    text(sl, xr3, yb + 13.4, wr, 2.8, [[B("Fig. 2 "), N("Cyanidin-3-O-glucoside (purple) in α-glucosidase (left) and α-amylase (right). "
         "Gold = catalytic residues, green = hydrogen-bond partners.")]], size=24)
    panel(sl, xr3, yb + 16.3, wr, 11.4, PALE, BERRY)
    text(sl, xr3 + 0.9, yb + 16.7, wr - 1.8, 10.6, [
        [B("The tightest binder is not the usable one", color=GOLD, size=30)],
        [N(f"γ-Oryzanol esters topped the ranking (cycloartenyl ferulate {cf['MGAM']['mean']:.2f} / {cf['HPA']['mean']:.2f} kcal/mol)".replace("-", "\u2212")
           + ", but with logP > 10 they will not dissolve in the gut lumen, and they sit in the bran's oil fraction.")],
        [B("Anthocyanins are the only soluble front-runners", color=BERRY, size=30)],
        [N(f"Cyanidin-3-O-glucoside ({c3g['MGAM']['mean']:.2f} / {c3g['HPA']['mean']:.2f}) and peonidin-3-O-glucoside "
           f"({p3g['MGAM']['mean']:.2f} / {p3g['HPA']['mean']:.2f}) ranked next".replace("-", "\u2212")
           + f", hydrogen-bonded the catalytic acids {res('ASP443')}/{res('ASP542')} and {res('ASP197')}/{res('GLU233')}, "
             "and were the only high-ranking compounds with logP < 0.7.")],
    ], size=25, space=0.4)

    # ============ Row 4: from screen to bench to product ============
    y = 84.2
    yb = heading(sl, X0, y, W, "FROM SCREEN TO PRODUCT", BERRY)
    boxes = [
        ("① Take this fraction", LEAF, [
            [N("The "), B("acidified aqueous extract"), N(" of fresh Riceberry bran — the anthocyanin-bearing fraction.")],
            [N("Not whole bran, not the oil fraction: the γ-oryzanol that ranks highest sits in the oil and will not dissolve.")],
            [N("Bran must be fresh; it turns rancid within weeks.", italic=True)]]),
        ("② Prove it on the bench", NAVY, [
            [B("Two-enzyme IC₅₀, not one."), N(" Run α-glucosidase and α-amylase assays side by side against acarbose.")],
            [N("Report a "), B("selectivity index"), N(" (IC₅₀ amylase ÷ IC₅₀ glucosidase) so potency and selectivity are measured, not inferred.")],
            [N("Quantify anthocyanin by pH-differential (520/700 nm) as cyanidin-3-glucoside equivalents.")]]),
        ("③ Build it into food", BERRY, [
            [N("Anthocyanins survive "), B("acid and short heat"), N(", not long baking. The prototype is a no-added-sugar jelly drink: pH 3.2–3.5, 80 °C for 1 min, hot-filled.")],
            [N("Three extract levels (20/30/40%), 9-point hedonic test, n ≥ 30.")],
            [N("Label claims stay factual: no health claim without FDA approval.", italic=True)]]),
    ]
    bw4 = (W - 2 * 1.0) / 3
    for i, (h, c, body) in enumerate(boxes):
        bx = X0 + i * (bw4 + 1.0)
        panel(sl, bx, yb, bw4, 14.6, SOFT, c)
        text(sl, bx + 0.8, yb + 0.6, bw4 - 1.6, 13.4, [[B(h, color=c, size=30)]] + body, size=25)

    # ============ Row 5: conclusion + QR ============
    y = 103.0
    yb = heading(sl, X0, y, W, "CONCLUSION")
    wq = 9.5
    text(sl, X0, yb, W - wq - 1.5, 11.5, [
        [N("Scoring binding "), B("together with solubility", color=BERRY),
         N(" changes which molecule a food project should chase. The highest-scoring constituents of Riceberry bran, the "),
         B("γ-oryzanol esters"), N(", cannot dissolve where the enzymes work; the "), B("anthocyanins"),
         N(" rank just below them and are the only front-runners that can.")],
        [N("This turns a rice-mill by-product into one defined fraction to extract and one assay that would settle it, instead of a vague claim about rice bran being healthy.")],
        [B("What this cannot tell you. ", color=NAVY), N("Docking ranks predicted binding, not inhibitory potency or selectivity; differences "
           "below ~1 kcal/mol sit inside the scoring error. Prior in vitro work established that these anthocyanins do inhibit both enzymes, and the "
           "two-enzyme IC₅₀ comparison above is what would rank them. No health claim is made here.")],
    ], size=27)
    qx = X0 + W - wq
    picture(sl, FIG / "qr.png", qx, yb, w=wq)
    text(sl, qx - 1.5, yb + 9.8, wq + 3.0, 2.4, [[B("Rotate the binding sites in 3D", size=23, color=NAVY)]], align=PP_ALIGN.CENTER)

    text(sl, X0, 118.0, W, 3.2, [[B("Structures "), N("PDB 2QMJ (Sim et al., J Mol Biol 2008) and 3BAJ (Maurus et al., Biochemistry 2008).  "),
        B("Software "), N("AutoDock Vina 1.2.7 (Eberhardt et al., JCIM 2021); Meeko; RDKit; PDBFixer. Ligands from PubChem.  "),
        B("Prior in vitro data "), N("Akkarachiyasit et al., Int J Mol Sci 2010;11:3387.  "),
        B("Acarbose tolerability "), N("US prescribing information.  All scripts and raw results via the QR link.")]], size=19)

    out = ROOT / "submission" / f"Poster_{A['file_name']}.pptx"; out.parent.mkdir(exist_ok=True)
    prs.save(out); print(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "")
